#!/opt/juno/venv/bin/python
"""Create, read and edit Excel workbooks (openpyxl).

  xlsx_tool.py create OUT.xlsx SPEC.json   sheets -> new workbook
  xlsx_tool.py dump FILE.xlsx              print every cell: coordinate, formula or value
  xlsx_tool.py set FILE.xlsx 'Sheet!A1' VALUE   set one cell (VALUE starting with = is a formula)
  xlsx_tool.py recalc FILE.xlsx            compute formulas with LibreOffice and print the results
"""
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_lib"))
from juno_skill import die, load_spec  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


def build_sheet(ws, spec):
    for row in spec["rows"]:
        ws.append(row)
    if spec.get("header", True) and spec["rows"]:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
    for column, fmt in spec.get("formats", {}).items():  # {"B": "#,##0.00"}
        for cell in ws[column][1:]:
            cell.number_format = fmt
    for i, column in enumerate(ws.columns, 1):
        widest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[get_column_letter(i)].width = min(max(widest + 2, 8), 60)
    if spec.get("freeze", "A2" if spec.get("header", True) else None):
        ws.freeze_panes = spec.get("freeze", "A2")


def cell_of(wb, ref):
    sheet, _, coordinate = ref.rpartition("!")
    return (wb[sheet] if sheet else wb.active)[coordinate]


def main(argv):
    cmd = argv[1] if len(argv) > 1 else ""
    if cmd == "create" and len(argv) == 4:
        wb = Workbook()
        wb.remove(wb.active)
        for spec in load_spec(argv[3])["sheets"]:
            build_sheet(wb.create_sheet(spec.get("name", "Sheet")), spec)
        wb.save(argv[2])
        print("wrote " + argv[2])
    elif cmd == "dump" and len(argv) == 3:
        for ws in load_workbook(argv[2]).worksheets:
            print("# " + ws.title)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        print(cell.coordinate + "\t" + str(cell.value))
    elif cmd == "set" and len(argv) == 5:
        wb = load_workbook(argv[2])
        value = argv[4]
        cell_of(wb, argv[3]).value = float(value) if value.replace(".", "", 1).lstrip("-").isdigit() else value
        wb.save(argv[2])
        print("set " + argv[3])
    elif cmd == "recalc" and len(argv) == 3:
        out = tempfile.mkdtemp(prefix="juno-recalc-")
        try:
            # The local computer image installs LibreOffice on first use: this call prints its status line to stderr and waits (no time limit), then it is there for good.
            subprocess.run(["soffice", "--version"], check=True, stdout=subprocess.DEVNULL, stderr=None)
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, argv[2]], check=True, timeout=120,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env={**os.environ, "HOME": out})
            for ws in load_workbook(os.path.join(out, os.path.basename(argv[2])), data_only=True).worksheets:
                print("# " + ws.title)
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            print(cell.coordinate + "\t" + str(cell.value))
        finally:
            shutil.rmtree(out, ignore_errors=True)
    else:
        die(__doc__)


main(sys.argv)
